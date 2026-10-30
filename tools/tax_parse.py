"""Reads tools/tax_source.xlsx (the Taxation study workbook) into a list of pages made of simple blocks."""
import os, re
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
DARK = {"FF9C1C1C", "FF1F3864", "FF00695C", "FFA5470B", "FF5B2C83", "FF375623", "FF7B3F00"}
PREFIX = {
    "bigidea": r"^THE BIG IDEA \(read this first\):\s*",
    "hook": r"^HOOK:\s*",
    "trap": r"^TRAP:\s*",
    "added": r"^ADDED BY ME \(not in your original; double-check\):\s*",
}

def fill_of(c):
    return c.fill.fgColor.rgb if c.fill is not None and c.fill.fill_type else None

def cell_info(c):
    col = c.font.color.rgb if c.font is not None and c.font.color is not None and isinstance(c.font.color.rgb, str) else None
    return dict(col=c.column, text=str(c.value).strip(), fill=fill_of(c), bold=bool(c.font.bold),
                red=col in ("FFC00000", "FF9C0006"), blue=col in ("FF1F4E9E", "FF0000FF"))

def parse_sheet(ws):
    blocks, table = [], None
    def close():
        nonlocal table
        if table:
            used = sorted(set(table["hc"]) | {k for row in table["rows"] for k in row})
            table["head"] = [table["hc"].get(k, "") for k in used]
            table["rows"] = [[row.get(k) for k in used] for row in table["rows"]]
            del table["hc"]
            blocks.append(table)
        table = None
    for r in range(1, ws.max_row + 1):
        cells = [cell_info(ws.cell(r, c)) for c in range(1, ws.max_column + 1) if ws.cell(r, c).value is not None and str(ws.cell(r, c).value).strip() != ""]
        if not cells:
            close(); continue
        first = cells[0]; f = first["fill"]; t = first["text"]
        if r == 1:
            blocks.append(dict(kind="title", text=t)); continue
        if f in DARK and len(cells) == 1:
            close(); blocks.append(dict(kind="heading", text=t)); continue
        if f == "FFD9E1F2":
            close()
            hc = {c["col"]: c["text"] for c in cells}
            table = dict(kind="table", hc=hc, rows=[]); continue
        if f == "FFFFF2CC" and len(cells) == 1:
            close()
            if re.match(PREFIX["bigidea"], t): blocks.append(dict(kind="bigidea", text=re.sub(PREFIX["bigidea"], "", t)))
            else: blocks.append(dict(kind="note", text=t))
            continue
        kinds = {"FFE2EFDA": "hook", "FFFCE4D6": "trap", "FFDDEBF7": "added", "FFFFE699": "fivehead", "FFFFF8DC": "five"}
        if f in kinds and len(cells) == 1:
            close(); k = kinds[f]
            blocks.append(dict(kind=k, text=re.sub(PREFIX[k], "", t) if k in PREFIX else t)); continue
        if table is not None:
            table["rows"].append({c["col"]: c for c in cells})
            continue
        if len(cells) == 1:
            if t.startswith("•"): blocks.append(dict(kind="bullet", text=t.lstrip("• ").strip(), cell=first))
            else: blocks.append(dict(kind="para", text=t, cell=first))
            continue
        # several cells with no header row: show as a plain two-column table
        if blocks and blocks[-1]["kind"] == "loose":
            blocks[-1]["rows"].append(cells)
        else:
            blocks.append(dict(kind="loose", rows=[cells]))
    close()
    return blocks

def load():
    wb = openpyxl.load_workbook(os.path.join(HERE, "tax_source.xlsx"))
    return [(ws.title, parse_sheet(ws)) for ws in wb.worksheets]

if __name__ == "__main__":
    import collections
    for name, blocks in load():
        cnt = collections.Counter(b["kind"] for b in blocks)
        print(f"{name:24s}", dict(cnt))
        for b in blocks:
            if b["kind"] == "loose":
                print("   LOOSE:", [[c["text"][:30] for c in r] for r in b["rows"]][:3])
            if b["kind"] == "table" and any(len([x for x in row if x]) == 0 for row in b["rows"]):
                print("   EMPTY ROW in table", b["head"])
